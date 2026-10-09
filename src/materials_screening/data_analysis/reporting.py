"""Fixed plotting, reporting and safe exports for data-analysis results."""

from __future__ import annotations

import csv
import io
import json
import math
from collections.abc import Sequence
from datetime import date
from pathlib import Path
from typing import Any, Literal

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # type: ignore[import-untyped]  # noqa: E402

from materials_screening.data_analysis.dataset_store import DatasetStore
from materials_screening.data_analysis.interpretation import (
    ScientificInterpretationService,
)
from materials_screening.data_analysis.models import DataAnalysisArtifact
from materials_screening.data_analysis.scientific import (
    ScientificAnalysisBrief,
    ScientificBriefService,
)
from materials_screening.data_analysis.service import DataAnalysisService

PlotType = Literal["histogram", "boxplot", "scatter", "line", "bar", "heatmap"]
ReportFormat = Literal["docx", "md", "json", "csv"]


class DataAnalysisReportingService:
    """Create bounded artifacts without accepting arbitrary plotting code."""

    def __init__(self, store: DatasetStore) -> None:
        self.store = store

    def create_plot(
        self,
        dataset_id: str,
        *,
        plot_type: PlotType,
        x: str | None = None,
        y: str | None = None,
        group_by: str | None = None,
        columns: Sequence[str] = (),
        title: str | None = None,
        x_label: str | None = None,
        y_label: str | None = None,
        correlation_method: Literal["pearson", "spearman"] = "pearson",
    ) -> DataAnalysisArtifact:
        frame = self.store.load_dataframe(dataset_id)
        _validate_label(title, "title")
        _validate_label(x_label, "x_label")
        _validate_label(y_label, "y_label")
        if len(frame.index) > 10_000 and plot_type in {"scatter", "line"}:
            raise ValueError("scatter and line plots support at most 10,000 rows")
        figure, axis = plt.subplots(figsize=(8, 5), dpi=150)
        try:
            if plot_type == "histogram":
                _histogram(axis, frame, _required_column(frame, x, "x"), group_by)
            elif plot_type == "boxplot":
                _boxplot(axis, frame, _required_column(frame, y, "y"), group_by)
            elif plot_type == "scatter":
                _scatter(
                    axis,
                    frame,
                    _required_column(frame, x, "x"),
                    _required_column(frame, y, "y"),
                    group_by,
                )
            elif plot_type == "line":
                _line(
                    axis,
                    frame,
                    _required_column(frame, x, "x"),
                    _required_column(frame, y, "y"),
                    group_by,
                )
            elif plot_type == "bar":
                _bar(
                    axis,
                    frame,
                    _required_column(frame, x, "x"),
                    _required_column(frame, y, "y"),
                    group_by,
                )
            elif plot_type == "heatmap":
                _heatmap(axis, frame, columns, correlation_method)
            else:
                raise ValueError(f"unsupported plot type: {plot_type}")
            axis.set_title(title or _default_title(plot_type))
            if x_label is not None:
                axis.set_xlabel(x_label)
            if y_label is not None:
                axis.set_ylabel(y_label)
            figure.tight_layout()
            output = io.BytesIO()
            figure.savefig(output, format="png", dpi=150, bbox_inches="tight")
            content = output.getvalue()
        finally:
            plt.close(figure)
        return self.store.save_artifact(
            content,
            artifact_type="analysis_plot",
            dataset_id=dataset_id,
            display_name=f"{plot_type}.png",
            file_extension="png",
            media_type="image/png",
        )

    def create_report(
        self,
        dataset_id: str,
        *,
        analysis_ids: Sequence[str],
        plot_artifact_ids: Sequence[str] = (),
        format: ReportFormat = "docx",
        scientific_brief: ScientificAnalysisBrief | None = None,
    ) -> DataAnalysisArtifact:
        checked_analysis_ids = _unique_ids(analysis_ids, "analysis_ids", 100)
        if not checked_analysis_ids:
            raise ValueError("report requires at least one analysis_id")
        analyses = [self.store.get_analysis(value) for value in checked_analysis_ids]
        if any(result.dataset_id != dataset_id for result in analyses):
            raise ValueError("all analyses must belong to the report dataset")
        checked_plots = _unique_ids(plot_artifact_ids, "plot_artifact_ids", 50)
        plots = [self.store.get_artifact(value) for value in checked_plots]
        if any(
            item.dataset_id != dataset_id or item.artifact_type != "analysis_plot"
            for item in plots
        ):
            raise ValueError("report plots must be analysis_plot artifacts for dataset")
        if format == "docx":
            if scientific_brief is None:
                raise ValueError(
                    "scientific_brief is required for a researcher-facing DOCX report"
                )
            ScientificBriefService.require_confirmed(scientific_brief)
            if scientific_brief.dataset_id != dataset_id:
                raise ValueError("scientific_brief belongs to a different dataset")
            content = _scientific_docx_report(
                self.store,
                dataset_id,
                analyses,
                plots,
                scientific_brief,
            )
            media_type = (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        elif format == "md":
            content = _markdown_report(dataset_id, analyses, plots).encode("utf-8")
            media_type = "text/markdown"
        elif format == "json":
            content = json.dumps(
                {
                    "dataset_id": dataset_id,
                    "analyses": [item.model_dump(mode="json") for item in analyses],
                    "plot_artifact_ids": list(checked_plots),
                },
                ensure_ascii=False,
                indent=2,
                allow_nan=False,
            ).encode("utf-8")
            media_type = "application/json"
        elif format == "csv":
            content = _analysis_csv(analyses)
            media_type = "text/csv"
        else:
            raise ValueError("report format must be docx, md, json or csv")
        return self.store.save_artifact(
            content,
            artifact_type="analysis_report",
            dataset_id=dataset_id,
            analysis_ids=checked_analysis_ids,
            display_name=f"analysis-report.{format}",
            file_extension=format,
            media_type=media_type,
        )

    def export_dataset(
        self, dataset_id: str, *, format: Literal["csv", "json"]
    ) -> DataAnalysisArtifact:
        frame = self.store.load_dataframe(dataset_id)
        if format == "csv":
            safe = _escape_formula_cells(frame)
            output = io.StringIO(newline="")
            safe.to_csv(output, index=False, lineterminator="\n")
            content = ("\ufeff" + output.getvalue()).encode("utf-8")
            media_type = "text/csv"
        elif format == "json":
            records = _json_records(frame)
            content = json.dumps(
                records, ensure_ascii=False, indent=2, allow_nan=False
            ).encode("utf-8")
            media_type = "application/json"
        else:
            raise ValueError("dataset export format must be csv or json")
        return self.store.save_artifact(
            content,
            artifact_type="dataset_file",
            dataset_id=dataset_id,
            display_name=f"dataset-export.{format}",
            file_extension=format,
            media_type=media_type,
        )


def _histogram(axis: Any, frame: pd.DataFrame, x: str, group_by: str | None) -> None:
    _numeric_column(frame, x)
    if group_by is None:
        axis.hist(_finite_series(frame[x]), bins=20, alpha=0.8)
    else:
        _require_existing(frame, group_by)
        for label, group in frame.groupby(group_by, dropna=False, sort=True):
            values = _finite_series(group[x])
            if not values.empty:
                axis.hist(values, bins=20, alpha=0.5, label=str(label))
        axis.legend()
    axis.set_xlabel(x)
    axis.set_ylabel("Count")


def _boxplot(axis: Any, frame: pd.DataFrame, y: str, group_by: str | None) -> None:
    _numeric_column(frame, y)
    if group_by is None:
        numeric_values = _finite_series(frame[y])
        _require_non_empty(numeric_values, y)
        axis.boxplot([numeric_values], tick_labels=[y])
    else:
        _require_existing(frame, group_by)
        grouped_values: list[pd.Series[Any]] = []
        labels: list[str] = []
        for label, group in frame.groupby(group_by, dropna=False, sort=True):
            numeric = _finite_series(group[y])
            if not numeric.empty:
                grouped_values.append(numeric)
                labels.append(str(label))
        if not grouped_values:
            raise ValueError("boxplot has no finite values")
        axis.boxplot(grouped_values, tick_labels=labels)
    axis.set_ylabel(y)


def _scatter(
    axis: Any, frame: pd.DataFrame, x: str, y: str, group_by: str | None
) -> None:
    pair = _finite_pair(frame, x, y)
    if group_by is None:
        axis.scatter(pair[x], pair[y], alpha=0.75)
    else:
        _require_existing(frame, group_by)
        pair = pair.join(frame[[group_by]], how="left")
        for label, group in pair.groupby(group_by, dropna=False, sort=True):
            axis.scatter(group[x], group[y], alpha=0.75, label=str(label))
        axis.legend()
    axis.set_xlabel(x)
    axis.set_ylabel(y)


def _line(
    axis: Any, frame: pd.DataFrame, x: str, y: str, group_by: str | None
) -> None:
    pair = _finite_pair(frame, x, y)
    if group_by is None:
        ordered = pair.sort_values(x)
        axis.plot(ordered[x], ordered[y], marker="o")
    else:
        _require_existing(frame, group_by)
        pair = pair.join(frame[[group_by]], how="left")
        for label, group in pair.groupby(group_by, dropna=False, sort=True):
            ordered = group.sort_values(x)
            axis.plot(ordered[x], ordered[y], marker="o", label=str(label))
        axis.legend()
    axis.set_xlabel(x)
    axis.set_ylabel(y)


def _bar(
    axis: Any, frame: pd.DataFrame, x: str, y: str, group_by: str | None
) -> None:
    _require_existing(frame, x)
    _numeric_column(frame, y)
    keys = [x] if group_by is None else [x, group_by]
    if group_by is not None:
        _require_existing(frame, group_by)
    valid = frame[keys + [y]].copy()
    valid[y] = pd.to_numeric(valid[y], errors="coerce")
    valid = valid[valid[y].map(_is_finite)]
    grouped = valid.groupby(keys, dropna=False, sort=True)[y].agg(["mean", "sem"])
    if grouped.empty:
        raise ValueError("bar plot has no finite values")
    labels = [
        " / ".join(map(str, key if isinstance(key, tuple) else (key,)))
        for key in grouped.index
    ]
    errors = grouped["sem"].fillna(0.0)
    axis.bar(labels, grouped["mean"], yerr=errors, capsize=3)
    axis.tick_params(axis="x", rotation=30)
    axis.set_ylabel(f"Mean {y}")


def _heatmap(
    axis: Any,
    frame: pd.DataFrame,
    columns: Sequence[str],
    method: str,
) -> None:
    selected = tuple(columns)
    if not 2 <= len(selected) <= 20 or len(selected) != len(set(selected)):
        raise ValueError("heatmap requires 2 to 20 unique numeric columns")
    for column in selected:
        _numeric_column(frame, column)
    if method not in {"pearson", "spearman"}:
        raise ValueError("heatmap method must be pearson or spearman")
    matrix = frame[list(selected)].corr(method=method)
    if matrix.isna().all().all():
        raise ValueError("heatmap correlation matrix has no finite values")
    image = axis.imshow(matrix, vmin=-1, vmax=1, cmap="coolwarm")
    axis.set_xticks(range(len(selected)), labels=selected, rotation=45, ha="right")
    axis.set_yticks(range(len(selected)), labels=selected)
    axis.figure.colorbar(image, ax=axis, label="Correlation")


def _scientific_docx_report(
    store: DatasetStore,
    dataset_id: str,
    analyses: Sequence[Any],
    plots: Sequence[DataAnalysisArtifact],
    brief: ScientificAnalysisBrief,
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "科研数据分析报告"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _style_runs(header, 9, "666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("第 ")
    _set_docx_run_font(footer_run, 9, "666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    end_run = footer.add_run(" 页")
    _set_docx_run_font(end_run, 9, "666666")

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("科研数据分析报告")
    _set_docx_run_font(title_run, 23, "0B2545", bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run(brief.research_question)
    _set_docx_run_font(subtitle_run, 14, "44546A")
    for label, value in (
        ("数据", store.get(dataset_id).display_name),
        ("领域", "材料科学" if brief.domain == "materials_science" else "通用科研"),
        ("分析日期", date.today().isoformat()),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        _set_docx_run_font(label_run, 10.5, "333333", bold=True)
        value_run = paragraph.add_run(value)
        _set_docx_run_font(value_run, 10.5, "333333")

    inspection = DataAnalysisService(store).inspect_dataset(dataset_id, limit=5)
    primary = next(
        (
            analysis
            for analysis in reversed(analyses)
            if analysis.analysis_type
            in {"statistical_test", "regression", "descriptive"}
        ),
        None,
    )
    if primary is not None:
        interpretation = ScientificInterpretationService(store).interpret(
            brief, primary
        )
        findings = [
            interpretation.direct_answer,
            f"证据强度：{interpretation.evidence_strength}。"
            f"{interpretation.evidence_explanation}",
            *interpretation.key_findings,
            interpretation.practical_significance,
        ]
        limitations = list(interpretation.limitations)
        recommendations = list(interpretation.recommendations)
    else:
        findings = _scientific_findings(analyses, brief)
        limitations = _scientific_limitations(analyses, inspection, brief)
        recommendations = _scientific_recommendations(
            analyses, inspection, brief
        )

    document.add_heading("一页式摘要", level=1)
    lead = document.add_paragraph()
    lead_run = lead.add_run(findings[0] if findings else "当前结果不足以形成主要结论。")
    _set_docx_run_font(lead_run, 11.5, "1F3A5F", bold=True)
    for item in findings[1:4]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
    if limitations:
        caution = document.add_paragraph()
        caution_run = caution.add_run(f"解读时最需要注意：{limitations[0]}")
        _set_docx_run_font(caution_run, 10.5, "7A5A00", bold=True)

    document.add_heading("1. 研究问题与分析边界", level=1)
    document.add_paragraph(f"研究问题：{brief.research_question}")
    document.add_paragraph(f"观测单位：{brief.observation_unit}")
    document.add_paragraph(
        "实验设计：" + _design_label(brief.design)
    )
    document.add_paragraph(
        "响应变量：" + "、".join(brief.response_variables)
    )
    if brief.group_variable:
        document.add_paragraph(f"分组/处理变量：{brief.group_variable}")
    if brief.hypothesis:
        document.add_paragraph(f"待检验假设：{brief.hypothesis}")
    document.add_paragraph(
        "本报告依据当前数据和已确认变量角色进行统计描述与推断。"
        "统计关联本身不构成因果证据。"
    )

    document.add_heading("2. 数据质量与可用性", level=1)
    document.add_paragraph(
        f"数据包含 {inspection.dataset.row_count} 条观测、"
        f"{inspection.dataset.column_count} 个字段。"
        f"共识别 {len(inspection.issues)} 项质量提示。"
    )
    if inspection.issues:
        for issue in inspection.issues[:10]:
            document.add_paragraph(issue.message, style="List Bullet")
    else:
        document.add_paragraph(
            "未发现缺失、重复、常量列或 IQR 异常值提示。"
        )

    document.add_heading("3. 主要分析结果", level=1)
    for index, analysis in enumerate(analyses, 1):
        document.add_heading(
            f"3.{index} {_analysis_heading(analysis)}", level=2
        )
        document.add_paragraph(_analysis_narrative(analysis, brief))
        _add_analysis_table(document, analysis)
        if analysis.warnings:
            for warning in analysis.warnings:
                document.add_paragraph(f"注意：{warning}", style="List Bullet")

    for index, plot in enumerate(plots, 1):
        path = store.resolve_artifact_path(plot.artifact_id)
        if path.is_file():
            document.add_picture(str(path), width=Inches(6.2))
            caption = document.add_paragraph(
                f"图 {index}  {Path(plot.display_name).stem}"
            )
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_runs(caption, 9.5, "555555")

    document.add_heading("4. 科研解释", level=1)
    for item in findings:
        document.add_paragraph(item)

    document.add_heading("5. 局限与不确定性", level=1)
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("6. 建议的下一步", level=1)
    for item in recommendations:
        document.add_paragraph(item, style="List Number")

    document.add_page_break()  # type: ignore[no-untyped-call]
    document.add_heading("附录：方法与可复现性记录", level=1)
    document.add_paragraph(
        "本附录保留方法和稳定记录标识，供复核使用；正文结论不依赖读者理解这些标识。"
    )
    audit_rows = [["分析类型", "方法", "分析记录"]]
    audit_rows.extend(
        [analysis.analysis_type, analysis.method, analysis.analysis_id]
        for analysis in analyses
    )
    table = document.add_table(rows=0, cols=3)
    for row_values in audit_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values, strict=True):
            cell.text = str(value)
    _format_docx_table(table, (2100, 2460, 4800))

    properties = document.core_properties
    properties.title = "科研数据分析报告"
    properties.subject = brief.research_question
    properties.author = "Materials Multi-Agent"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _scientific_findings(
    analyses: Sequence[Any], brief: ScientificAnalysisBrief
) -> list[str]:
    findings: list[str] = []
    for analysis in analyses:
        if analysis.analysis_type == "statistical_test":
            significant = analysis.summary.get("significant") is True
            response = analysis.parameters.get("response_column") or "目标指标"
            p_value = _report_value(analysis.summary.get("p_value"))
            groups = analysis.summary.get("groups", [])
            means = "；".join(
                f"{item.get('label')} 组均值 {_report_value(item.get('mean'))}"
                for item in groups
            )
            if significant:
                conclusion = f"{response} 在组间呈现统计学差异（p={p_value}）"
            else:
                conclusion = (
                    f"当前数据没有提供足够证据证明 {response} 存在组间差异"
                    f"（p={p_value}）"
                )
            findings.append(f"{conclusion}。{means}。")
            significant_pairs = [
                item
                for item in analysis.summary.get("post_hoc", [])
                if item.get("significant") is True
            ]
            if significant_pairs:
                pair_text = "；".join(
                    f"{item.get('group_a')}–{item.get('group_b')}"
                    f"（校正后 p={_report_value(item.get('adjusted_p_value'))}）"
                    for item in significant_pairs[:6]
                )
                findings.append(
                    f"Holm 校正后的事后比较提示以下组对存在差异：{pair_text}。"
                )
            threshold = brief.practical_thresholds.get(str(response))
            difference = analysis.summary.get("mean_difference")
            if threshold is not None and difference is not None:
                practical = abs(float(difference)) >= threshold
                findings.append(
                    f"观察到的均值差为 {_report_value(difference)}，"
                    f"{'达到' if practical else '未达到'}预设的实际意义阈值 "
                    f"{_report_value(threshold)}。"
                )
        elif analysis.analysis_type == "correlation":
            pairs = [
                item
                for item in analysis.summary.get("pairs", [])
                if item.get("x") != item.get("y")
                and item.get("coefficient") is not None
            ]
            if pairs:
                strongest = max(
                    pairs, key=lambda item: abs(float(item["coefficient"]))
                )
                findings.append(
                    f"最强字段关联出现在 {strongest['x']} 与 {strongest['y']} "
                    f"之间（r={_report_value(strongest['coefficient'])}，"
                    f"n={strongest['n']}）；该结果不能单独解释为因果作用。"
                )
    if not findings:
        findings.append("当前分析以数据描述为主，尚不足以回答明确的推断性研究问题。")
    return findings


def _scientific_limitations(
    analyses: Sequence[Any], inspection: Any, brief: ScientificAnalysisBrief
) -> list[str]:
    limitations: list[str] = []
    if inspection.dataset.row_count < 30:
        limitations.append("总样本量少于 30，估计区间和显著性结论可能不稳定。")
    if inspection.issues:
        limitations.append(
            "数据存在质量提示；缺失、重复或异常候选可能影响估计，需回查实验记录。"
        )
    if not brief.practical_thresholds:
        limitations.append(
            "尚未设置实际意义阈值，因此只能判断统计证据，不能代替科研价值判断。"
        )
    if any(analysis.warnings for analysis in analyses):
        limitations.append("部分分析存在方法前提警告，结论应结合稳健方法复核。")
    limitations.append(
        "本分析基于观测到的字段，未控制的批次、仪器、制样或环境因素仍可能造成混杂。"
    )
    return limitations


def _scientific_recommendations(
    analyses: Sequence[Any], inspection: Any, brief: ScientificAnalysisBrief
) -> list[str]:
    recommendations = [
        "结合实验记录确认分组、样品独立性、单位和缺失值含义。",
    ]
    if inspection.issues:
        recommendations.append("逐条复核质量提示，并保留清洗前后的数据版本。")
    if not brief.practical_thresholds:
        recommendations.append(
            "在追加检验前，由课题负责人给出具有学科意义的最小差异阈值。"
        )
    if any(
        analysis.analysis_type == "statistical_test"
        and analysis.method == "anova"
        and analysis.summary.get("significant") is True
        for analysis in analyses
    ):
        recommendations.append(
            "ANOVA 显著时执行带多重比较校正的事后分析，确认具体差异组对。"
        )
    recommendations.append(
        "在独立批次或新增样品上复现主要结果，再讨论可能的材料机理。"
    )
    return recommendations


def _analysis_heading(analysis: Any) -> str:
    return {
        "descriptive": "数据分布与描述统计",
        "correlation": "字段关联分析",
        "statistical_test": "组间差异检验",
        "quality": "数据质量评估",
        "regression": "连续变量关系分析",
    }.get(analysis.analysis_type, "分析结果")


def _analysis_narrative(analysis: Any, brief: ScientificAnalysisBrief) -> str:
    if analysis.analysis_type == "statistical_test":
        return _scientific_findings((analysis,), brief)[0]
    if analysis.analysis_type == "correlation":
        return _scientific_findings((analysis,), brief)[0]
    if analysis.analysis_type == "regression":
        if analysis.summary.get("significant") is True:
            return (
                "当前数据支持响应指标与预测因素之间存在线性关联；"
                "这不等同于因果作用。"
            )
        return "当前数据没有提供足够证据确认线性关联；这不等同于证明不存在关联。"
    if analysis.analysis_type == "descriptive":
        count = len(analysis.summary.get("statistics", []))
        return (
            f"本节汇总了 {count} 组字段统计，用于理解中心趋势、"
            "离散程度和样本可用性。"
        )
    return "本节记录与研究问题相关的确定性分析结果。"


def _add_analysis_table(document: Any, analysis: Any) -> None:
    if analysis.analysis_type == "statistical_test":
        rows = [["组别", "样本数", "均值"]]
        rows.extend(
            [
                str(item.get("label", "—")),
                str(item.get("n", "—")),
                _report_value(item.get("mean")),
            ]
            for item in analysis.summary.get("groups", [])
        )
        table = document.add_table(rows=0, cols=3)
        for values in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
        _format_docx_table(table, (3120, 2160, 4080))
    elif analysis.analysis_type == "descriptive":
        statistics = analysis.summary.get("statistics", [])[:12]
        if not statistics:
            return
        table = document.add_table(rows=0, cols=4)
        rows = [["字段", "样本数", "均值", "标准差"]]
        rows.extend(
            [
                str(item.get("column", "—")),
                str(item.get("count", "—")),
                _report_value(item.get("mean")),
                _report_value(item.get("std")),
            ]
            for item in statistics
        )
        for values in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
        _format_docx_table(table, (3300, 1800, 2130, 2130))


def _format_docx_table(table: Any, widths: tuple[int, ...]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row_index, row in enumerate(table.rows):
        for cell, value in zip(row.cells, widths, strict=True):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(value))
            cell_width.set(qn("w:type"), "dxa")
            margins = OxmlElement("w:tcMar")
            for side, amount in (
                ("top", 80),
                ("bottom", 80),
                ("start", 120),
                ("end", 120),
            ):
                element = OxmlElement(f"w:{side}")
                element.set(qn("w:w"), str(amount))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                _style_runs(
                    paragraph,
                    9.5,
                    "333333",
                    bold=row_index == 0,
                )
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                cell._tc.get_or_add_tcPr().append(shading)
    table.style = "Table Grid"
    for row in table.rows:
        row.height = None
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string("1F3A5F")


def _style_runs(
    paragraph: Any,
    size: float,
    color: str,
    *,
    bold: bool | None = None,
) -> None:
    for run in paragraph.runs:
        _set_docx_run_font(run, size, color, bold=bold)


def _set_docx_run_font(
    run: Any,
    size: float,
    color: str,
    *,
    bold: bool | None = None,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _design_label(value: str) -> str:
    return {
        "exploratory": "探索性分析",
        "independent_groups": "独立组设计",
        "paired": "配对设计",
        "repeated_measures": "重复测量设计",
        "continuous_relationship": "连续变量关系分析",
    }.get(value, value)


def _report_value(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.4g}"
    return str(value)


def _markdown_report(
    dataset_id: str, analyses: Sequence[Any], plots: Sequence[DataAnalysisArtifact]
) -> str:
    lines = [
        "# 数据分析报告",
        "",
        f"- 数据集：`{dataset_id}`",
        f"- 分析数量：{len(analyses)}",
        f"- 图表数量：{len(plots)}",
        "",
        "## 分析结果",
    ]
    for index, result in enumerate(analyses, 1):
        lines.extend(
            (
                "",
                f"### {index}. {result.method}",
                "",
                f"- Analysis ID：`{result.analysis_id}`",
                f"- 类型：`{result.analysis_type}`",
                f"- Evidence ID：`{result.evidence_id}`",
                "",
                "```json",
                json.dumps(result.summary, ensure_ascii=False, indent=2),
                "```",
            )
        )
        if result.warnings:
            lines.extend(("", "警告："))
            lines.extend(f"- {warning}" for warning in result.warnings)
    if plots:
        lines.extend(("", "## 图表 Artifact", ""))
        lines.extend(f"- `{item.artifact_id}`：{item.display_name}" for item in plots)
    lines.extend(
        (
            "",
            "## 解释边界",
            "",
            "统计关系不自动代表因果关系；结论应结合实验设计、样本量和前提检查。",
        )
    )
    return "\n".join(lines) + "\n"


def _analysis_csv(analyses: Sequence[Any]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(
        output,
        fieldnames=[
            "analysis_id",
            "analysis_type",
            "method",
            "parameters_json",
            "summary_json",
            "warnings",
        ],
        lineterminator="\n",
    )
    writer.writeheader()
    for result in analyses:
        writer.writerow(
            {
                "analysis_id": _safe_csv_cell(result.analysis_id),
                "analysis_type": _safe_csv_cell(result.analysis_type),
                "method": _safe_csv_cell(result.method),
                "parameters_json": _safe_csv_cell(
                    json.dumps(result.parameters, ensure_ascii=False)
                ),
                "summary_json": _safe_csv_cell(
                    json.dumps(result.summary, ensure_ascii=False)
                ),
                "warnings": _safe_csv_cell(" | ".join(result.warnings)),
            }
        )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def _escape_formula_cells(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.select_dtypes(include=["object", "string"]).columns:
        result[column] = result[column].map(
            lambda value: _safe_csv_cell(value) if isinstance(value, str) else value
        )
    return result


def _safe_csv_cell(value: Any) -> str:
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _json_records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for row in frame.to_dict(orient="records"):
        converted: dict[str, Any] = {}
        for key, value in row.items():
            if value is None or bool(pd.isna(value)):
                converted[str(key)] = None
            elif hasattr(value, "item"):
                converted[str(key)] = value.item()
            elif hasattr(value, "isoformat"):
                converted[str(key)] = value.isoformat()
            elif isinstance(value, float) and not math.isfinite(value):
                converted[str(key)] = None
            else:
                converted[str(key)] = value
        records.append(converted)
    return records


def _finite_pair(frame: pd.DataFrame, x: str, y: str) -> pd.DataFrame:
    _numeric_column(frame, x)
    _numeric_column(frame, y)
    pair = frame[[x, y]].copy()
    pair[x] = pd.to_numeric(pair[x], errors="coerce")
    pair[y] = pd.to_numeric(pair[y], errors="coerce")
    pair = pair[pair[x].map(_is_finite) & pair[y].map(_is_finite)]
    if pair.empty:
        raise ValueError("plot has no complete finite x/y pairs")
    return pair


def _finite_series(series: pd.Series[Any]) -> pd.Series[Any]:
    numeric = pd.to_numeric(series, errors="coerce")
    return numeric[numeric.map(_is_finite)]


def _is_finite(value: Any) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _numeric_column(frame: pd.DataFrame, column: str) -> None:
    _require_existing(frame, column)
    original_count = int(frame[column].notna().sum())
    numeric_count = int(pd.to_numeric(frame[column], errors="coerce").notna().sum())
    if original_count != numeric_count:
        raise ValueError(f"plot column {column!r} must be numeric")


def _required_column(
    frame: pd.DataFrame, value: str | None, parameter_name: str
) -> str:
    if value is None:
        raise ValueError(f"plot requires {parameter_name}")
    _require_existing(frame, value)
    return value


def _require_existing(frame: pd.DataFrame, column: str) -> None:
    if column not in frame.columns:
        raise ValueError(f"unknown plot column: {column!r}")


def _require_non_empty(series: pd.Series[Any], column: str) -> None:
    if series.empty:
        raise ValueError(f"column {column!r} has no finite values")


def _validate_label(value: str | None, name: str) -> None:
    if value is not None and (not value.strip() or len(value) > 256):
        raise ValueError(f"{name} must be non-blank and at most 256 characters")


def _unique_ids(values: Sequence[str], name: str, limit: int) -> tuple[str, ...]:
    checked = tuple(values)
    if len(checked) > limit or len(checked) != len(set(checked)):
        raise ValueError(f"{name} must be unique and contain at most {limit} values")
    return checked


def _default_title(plot_type: str) -> str:
    return plot_type.replace("_", " ").title()
