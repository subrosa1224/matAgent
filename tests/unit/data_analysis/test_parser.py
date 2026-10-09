from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from materials_screening.data_analysis.parser import (
    TabularDataParser,
    schema_fingerprint,
)


def test_parser_accepts_utf8_bom_csv_and_preserves_general_columns(
    tmp_path: Path,
) -> None:
    path = tmp_path / "实验.csv"
    path.write_text(
        "sample_id,温度,conductivity\nA,800,1.2\nB,900,2.4\n",
        encoding="utf-8-sig",
    )

    frame = TabularDataParser().parse(path)

    assert list(frame.columns) == ["sample_id", "温度", "conductivity"]
    assert frame.shape == (2, 3)
    assert schema_fingerprint(frame).startswith("sha256:")


def test_parser_accepts_flat_json_object_array(tmp_path: Path) -> None:
    path = tmp_path / "data.json"
    path.write_text(
        json.dumps(
            [
                {"sample": "A", "value": 1.0},
                {"sample": "B", "value": None},
            ]
        ),
        encoding="utf-8",
    )

    frame = TabularDataParser().parse(path)

    assert frame.shape == (2, 2)
    assert list(frame["sample"]) == ["A", "B"]


def test_parser_accepts_first_xlsx_sheet_and_applies_shape_limits(
    tmp_path: Path,
) -> None:
    path = tmp_path / "experiment.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(["sample", "temperature", "value"])
    sheet.append(["A", 800, 1.2])
    sheet.append(["B", 900, 2.4])
    ignored = workbook.create_sheet("ignored")
    ignored.append(["secret"])
    ignored.append([999])
    workbook.save(path)

    frame = TabularDataParser().parse(path)

    assert frame.shape == (2, 3)
    assert list(frame.columns) == ["sample", "temperature", "value"]
    assert "secret" not in frame.columns
    with pytest.raises(ValueError, match="row limit"):
        TabularDataParser(max_rows=1).parse(path)


def test_parser_rejects_xlsx_duplicate_headers(tmp_path: Path) -> None:
    path = tmp_path / "duplicate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["value", "value"])
    sheet.append([1, 2])
    workbook.save(path)

    with pytest.raises(ValueError, match="must be unique"):
        TabularDataParser().parse(path)


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("empty.csv", "", "empty"),
        ("header.csv", "a,b\n", "at least one data row"),
        ("duplicate.csv", "a,a\n1,2\n", "must be unique"),
        ("blank.csv", "a,\n1,2\n", "must not be blank"),
        ("bad.json", "{oops", "invalid JSON"),
        ("object.json", '{"a": 1}', "top-level array"),
        ("empty.json", "[]", "at least one row"),
        ("scalar.json", "[1]", "must be an object"),
        ("nested.json", '[{"a": {"b": 1}}]', "contains nested data"),
    ],
)
def test_parser_rejects_invalid_or_unsupported_shapes(
    tmp_path: Path, name: str, content: str, message: str
) -> None:
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")

    with pytest.raises(ValueError, match=message):
        TabularDataParser().parse(path)


def test_parser_rejects_non_utf8_csv(tmp_path: Path) -> None:
    path = tmp_path / "gbk.csv"
    path.write_bytes("名称,值\n样品,1\n".encode("gbk"))

    with pytest.raises(ValueError, match="UTF-8"):
        TabularDataParser().parse(path)


def test_parser_enforces_configured_resource_limits(tmp_path: Path) -> None:
    path = tmp_path / "large.csv"
    path.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")

    with pytest.raises(ValueError, match="row limit"):
        TabularDataParser(max_rows=1).parse(path)
    with pytest.raises(ValueError, match="column limit"):
        TabularDataParser(max_columns=1).parse(path)
    with pytest.raises(ValueError, match="cell limit"):
        TabularDataParser(max_cells=3).parse(path)
    with pytest.raises(ValueError, match="byte limit"):
        TabularDataParser(max_file_bytes=2).parse(path)


def test_parser_rejects_oversized_text_cell(tmp_path: Path) -> None:
    path = tmp_path / "cell.csv"
    path.write_text("text\nabcdef\n", encoding="utf-8")

    with pytest.raises(ValueError, match="longer than 5"):
        TabularDataParser(max_cell_characters=5).parse(path)


def test_parser_rejects_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsm"
    path.write_bytes(b"not a supported workbook")

    with pytest.raises(ValueError, match="unsupported dataset format"):
        TabularDataParser().parse(path)
